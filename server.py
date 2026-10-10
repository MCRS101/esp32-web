import os
import secrets

from io import BytesIO

from flask import (
    Flask,
    request,
    jsonify,
    render_template,
    send_file
)

from database import get_connection, init_database

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from datetime import datetime
from flask import request, jsonify
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from reportlab.platypus import Image as RLImage

from flask import session
from flask_bcrypt import Bcrypt
from functools import wraps

# =========================================================
# FLASK
# =========================================================

app = Flask(__name__)

app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY")

if not app.config["SECRET_KEY"]:
    raise RuntimeError("Please set SECRET_KEY environment variable")

app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SECURE"] = True  # ใช้ HTTPS
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

bcrypt = Bcrypt(app)

# =========================================================
# INITIALIZE DATABASE
# =========================================================
try:
    init_database()
    app.logger.info("Database initialization completed")

except Exception:
    app.logger.exception("MYSQL DATABASE INITIALIZATION FAILED")
    raise

# =========================================================
# HOME
# =========================================================

@app.route("/")
def index():

    return render_template("index.html")


# =========================================================
# AUTHENTICATION
# =========================================================

def login_required(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            return jsonify({
                "success": False,
                "message": "Please login first"
            }), 401

        return func(*args, **kwargs)

    return wrapper


# =========================================================
# DEVICE ACCESS CONTROL
# =========================================================

def device_access_required(func):
    @wraps(func)

    def wrapper(device_id, *args, **kwargs):

        # ต้อง Login ก่อน
        if "user_id" not in session:
            return jsonify({
                "success": False,
                "message": "Please login first"
            }), 401

        conn = get_connection()
        cursor = conn.cursor()

        try:
            cursor.execute("""
                SELECT 1
                FROM user_devices
                WHERE user_id = %s
                  AND device_id = %s
                LIMIT 1
            """, (session["user_id"], device_id))

            if cursor.fetchone() is None:
                return jsonify({
                    "success": False,
                    "message": "You do not have access to this device"
                }), 403

        finally:
            cursor.close()
            conn.close()

        return func(device_id, *args, **kwargs)

    return wrapper


# REGISTER
@app.route("/api/auth/register", methods=["POST"])
def register_user():
    data = request.get_json(silent=True) or {}

    username = data.get("username", "").strip()
    email = data.get("email", "").strip().lower()
    password = data.get("password", "")

    if not username or not email or not password:
        return jsonify({
            "success": False,
            "message": "username, email and password are required"
        }), 400

    if len(username) > 50 or len(email) > 255:
        return jsonify({
            "success": False,
            "message": "Username or email is too long"
        }), 400

    if len(password) < 8:
        return jsonify({
            "success": False,
            "message": "Password must be at least 8 characters"
        }), 400

    if len(password.encode("utf-8")) > 72:
        return jsonify({
            "success": False,
            "message": "Password is too long"
        }), 400

    password_hash = bcrypt.generate_password_hash(
        password
    ).decode("utf-8")

    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            INSERT INTO users (username, email, password_hash)
            VALUES (%s, %s, %s)
        """, (username, email, password_hash))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Registration successful"
        }), 201

    except Exception as e:
        conn.rollback()

        # Duplicate username/email
        if getattr(e, "errno", None) == 1062:
            return jsonify({
                "success": False,
                "message": "Username or email already exists"
            }), 409

        app.logger.exception("REGISTER ERROR")
        return jsonify({
            "success": False,
            "message": "Registration failed"
        }), 500

    finally:
        cursor.close()
        conn.close()


# LOGIN
@app.route("/api/auth/login", methods=["POST"])
def login_user():
    data = request.get_json(silent=True) or {}

    username = data.get("username", "").strip()
    password = data.get("password", "")

    if not username or not password:
        return jsonify({
            "success": False,
            "message": "Username and password are required"
        }), 400

    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT id, username, email, password_hash
            FROM users
            WHERE username = %s OR email = %s
            LIMIT 1
        """, (username, username.lower()))

        user = cursor.fetchone()

        if (
            not user
            or not bcrypt.check_password_hash(
                user["password_hash"], password
            )
        ):
            return jsonify({
                "success": False,
                "message": "Invalid username or password"
            }), 401

        # Prevent session fixation
        session.clear()
        session["user_id"] = user["id"]

        return jsonify({
            "success": True,
            "message": "Login successful",
            "user": {
                "id": user["id"],
                "username": user["username"],
                "email": user["email"]
            }
        })

    finally:
        cursor.close()
        conn.close()


# CURRENT USER
@app.route("/api/auth/me", methods=["GET"])
@login_required
def current_user():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT id, username, email, created_at
            FROM users
            WHERE id = %s
        """, (session["user_id"],))

        user = cursor.fetchone()

        if not user:
            session.clear()
            return jsonify({
                "success": False,
                "message": "User not found"
            }), 404

        if user.get("created_at"):
            user["created_at"] = user["created_at"].isoformat()

        return jsonify({
            "success": True,
            "user": user
        })

    finally:
        cursor.close()
        conn.close()


# LOGOUT
@app.route("/api/auth/logout", methods=["POST"])
def logout_user():
    session.clear()

    return jsonify({
        "success": True,
        "message": "Logout successful"
    })


# =========================================================
# USER DEVICES
# =========================================================

# ตรวจสอบว่าอุปกรณ์เป็นของผู้ใช้ที่ Login อยู่หรือไม่
def user_owns_device(cursor, user_id, device_id):
    cursor.execute("""
        SELECT 1
        FROM user_devices
        WHERE user_id = %s AND device_id = %s
        LIMIT 1
    """, (user_id, device_id))

    return cursor.fetchone() is not None


# เพิ่มอุปกรณ์เข้าบัญชี
@app.route("/api/user/devices", methods=["POST"])
@login_required
def add_user_device():
    data = request.get_json(silent=True) or {}

    device_id = str(data.get("device_id", "")).strip()
    device_token = str(data.get("device_token", "")).strip()

    if not device_id or not device_token:
        return jsonify({
            "success": False,
            "message": "device_id and device_token are required"
        }), 400

    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        # ตรวจสอบว่า Device มีอยู่จริงและ Token ตรงกัน
        cursor.execute("""
            SELECT device_id
            FROM devices
            WHERE device_id = %s AND device_token = %s
            LIMIT 1
        """, (device_id, device_token))

        device = cursor.fetchone()

        if not device:
            return jsonify({
                "success": False,
                "message": "Invalid device ID or token"
            }), 403

        # ถ้ามีผู้ใช้ผูกอุปกรณ์นี้แล้ว จะไม่อนุญาตให้แย่งอุปกรณ์
        cursor.execute("""
            SELECT user_id
            FROM user_devices
            WHERE device_id = %s
            LIMIT 1
        """, (device_id,))

        owner = cursor.fetchone()

        if owner:
            if owner["user_id"] == session["user_id"]:
                return jsonify({
                    "success": True,
                    "message": "Device already added",
                    "device_id": device_id
                })

            return jsonify({
                "success": False,
                "message": "Device is already linked to another account"
            }), 409

        cursor.execute("""
            INSERT INTO user_devices (user_id, device_id)
            VALUES (%s, %s)
        """, (session["user_id"], device_id))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Device added successfully",
            "device_id": device_id
        }), 201

    except Exception:
        conn.rollback()
        app.logger.exception("ADD USER DEVICE ERROR")

        return jsonify({
            "success": False,
            "message": "Unable to add device"
        }), 500

    finally:
        cursor.close()
        conn.close()


# แสดงเฉพาะอุปกรณ์ของผู้ใช้ที่ Login อยู่
@app.route("/api/user/devices", methods=["GET"])
@login_required
def get_user_devices():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                d.device_id,
                d.status,
                d.ip_address,
                d.last_seen,
                ud.added_at
            FROM user_devices AS ud
            JOIN devices AS d
                ON d.device_id = ud.device_id
            WHERE ud.user_id = %s
            ORDER BY ud.added_at DESC
        """, (session["user_id"],))

        rows = cursor.fetchall()

        for row in rows:
            for key in ("last_seen", "added_at"):
                if row.get(key):
                    row[key] = row[key].isoformat()

        return jsonify({
            "success": True,
            "count": len(rows),
            "devices": rows
        })

    finally:
        cursor.close()
        conn.close()



# =========================================================
# REGISTER ESP32
# =========================================================

@app.route("/api/device/register", methods=["POST"])
def register_device():

    data = request.get_json(silent=True)

    if not data:

        return jsonify({
            "success": False,
            "message": "Invalid JSON"
        }), 400


    device_id = data.get("device_id")

    if not device_id:

        return jsonify({
            "success": False,
            "message": "device_id is required"
        }), 400


    conn = get_connection()
    cursor = conn.cursor(dictionary=True)


    try:

        # -----------------------------------------
        # ตรวจสอบ Device
        # -----------------------------------------

        cursor.execute(
            """
            SELECT *
            FROM devices
            WHERE device_id = %s
            """,
            (device_id,)
        )

        device = cursor.fetchone()


        # -----------------------------------------
        # Device มีอยู่แล้ว
        # -----------------------------------------

        if device:

            cursor.execute(
                """
                UPDATE devices

                SET status = 'online',

                    ip_address = %s,

                    last_seen = CURRENT_TIMESTAMP

                WHERE device_id = %s
                """,
                (
                    request.remote_addr,
                    device_id
                )
            )

            conn.commit()


            return jsonify({
                "success": True,
                "message": "Device already registered",
                "device_id": device_id
            })


        # -----------------------------------------
        # Device ใหม่
        # -----------------------------------------

        device_token = secrets.token_hex(32)


        cursor.execute(
            """
            INSERT INTO devices
            (
                device_id,
                device_token,
                status,
                ip_address,
                last_seen
            )

            VALUES
            (
                %s,
                %s,
                'online',
                %s,
                CURRENT_TIMESTAMP
            )
            """,
            (
                device_id,
                device_token,
                request.remote_addr
            )
        )


        conn.commit()


        return jsonify({

            "success": True,

            "message": "Device registered",

            "device_id": device_id,

            "device_token": device_token

        })


    except Exception as e:

        conn.rollback()

        print("REGISTER ERROR:", e)


        return jsonify({

            "success": False,

            "message": "Database error"

        }), 500


    finally:

        cursor.close()
        conn.close()


# =========================================================
# RECEIVE SENSOR DATA
# =========================================================

@app.route("/api/device/data", methods=["POST"])
def receive_sensor_data():

    data = request.get_json(silent=True)


    if not data:

        return jsonify({

            "success": False,

            "message": "Invalid JSON"

        }), 400


    device_id = data.get("device_id")


    if not device_id:

        return jsonify({

            "success": False,

            "message": "device_id is required"

        }), 400


    # =====================================================
    # SENSOR VALUES
    # =====================================================

    try:

        accel_x = float(
            data.get("accel_x", 0)
        )

        accel_y = float(
            data.get("accel_y", 0)
        )

        accel_z = float(
            data.get("accel_z", 0)
        )


        pga = float(
            data.get("pga", 0)
        )

        peak_pga = float(
            data.get("peak_pga", 0)
        )

        avg_pga = float(
            data.get("avg_pga", 0)
        )


        pendulum = float(
            data.get("pendulum", 0)
        )


        level = data.get(
            "level",
            "LOW"
        )


        direction = data.get(
            "direction",
            "-"
        )


        estimated_ml = float(
            data.get(
                "estimated_ml",
                0
            )
        )


    except (ValueError, TypeError):

        return jsonify({

            "success": False,

            "message": "Invalid sensor data"

        }), 400


    # =====================================================
    # MYSQL
    # =====================================================

    conn = get_connection()
    cursor = conn.cursor()


    try:

        # -----------------------------------------
        # INSERT SENSOR DATA
        # -----------------------------------------

        cursor.execute(
            """
            INSERT INTO sensor_data
            (
                device_id,

                accel_x,
                accel_y,
                accel_z,

                pga,
                peak_pga,
                avg_pga,

                pendulum,

                level,

                direction,

                estimated_ml
            )

            VALUES
            (
                %s,

                %s,
                %s,
                %s,

                %s,
                %s,
                %s,

                %s,

                %s,

                %s,

                %s
            )
            """,
            (
                device_id,

                accel_x,
                accel_y,
                accel_z,

                pga,
                peak_pga,
                avg_pga,

                pendulum,

                level,

                direction,

                estimated_ml
            )
        )


        # -----------------------------------------
        # UPDATE DEVICE
        # -----------------------------------------

        cursor.execute(
            """
            UPDATE devices

            SET status = 'online',

                ip_address = %s,

                last_seen = CURRENT_TIMESTAMP

            WHERE device_id = %s
            """,
            (
                request.remote_addr,
                device_id
            )
        )


        conn.commit()


        # -----------------------------------------
        # RESPONSE
        # -----------------------------------------

        return jsonify({

            "success": True,

            "message": "Sensor data received",

            "alert": level in [
                "HIGH",
                "SEVERE"
            ]

        })


    except Exception as e:

        conn.rollback()

        print("SENSOR DATABASE ERROR:", e)


        return jsonify({

            "success": False,

            "message": "Database error"

        }), 500


    finally:

        cursor.close()
        conn.close()


# =========================================================
# GET LATEST SENSOR DATA
# =========================================================

@app.route(
    "/api/device/<device_id>/latest",
    methods=["GET"]
)
@login_required
@device_access_required
def get_latest(device_id):

    conn = get_connection()

    cursor = conn.cursor(
        dictionary=True
    )


    try:

        cursor.execute(
            """
            SELECT *

            FROM sensor_data

            WHERE device_id = %s

            ORDER BY id DESC

            LIMIT 1
            """,
            (device_id,)
        )


        row = cursor.fetchone()


        if not row:

            return jsonify({

                "success": False,

                "message": "No data"

            }), 404


        return jsonify({

            "success": True,

            "data": row

        })


    finally:

        cursor.close()
        conn.close()


# =========================================================
# GET SENSOR HISTORY
# =========================================================

@app.route(
    "/api/device/<device_id>/history",
    methods=["GET"]
)
@login_required
@device_access_required
def get_history(device_id):

    conn = get_connection()

    cursor = conn.cursor(
        dictionary=True
    )

    try:

        cursor.execute(
            """
            SELECT
                id,
                device_id,
                timestamp,
                accel_x,
                accel_y,
                accel_z,
                pga,
                peak_pga,
                avg_pga,
                pendulum,
                level,
                direction,
                estimated_ml

            FROM sensor_data

            WHERE device_id = %s

            ORDER BY id DESC
            """,
            (
                device_id,
            )
        )

        rows = cursor.fetchall()


        # -----------------------------------------
        # แปลงวันที่ให้ JavaScript อ่านง่าย
        # -----------------------------------------

        for row in rows:

            if row.get("timestamp"):

                row["timestamp"] = (
                    row["timestamp"].isoformat()
                )


        return jsonify({

            "success": True,

            "count": len(rows),

            "data": rows

        })


    except Exception as e:

        print("HISTORY ERROR:", e)

        return jsonify({

            "success": False,

            "message": "Database error"

        }), 500


    finally:

        cursor.close()
        conn.close()

@app.route("/api/device/<device_id>/graph", methods=["GET"])
@login_required
@device_access_required
def get_graph_data(device_id):
    """
    ดึงข้อมูลสำหรับกราฟรอบเวลาที่ผู้ใช้เลือก
    เช่น ก่อนหน้า 30 รายการ + หลัง 30 รายการ
    """

    selected_id = request.args.get("id", type=int)

    if not selected_id:
        return jsonify({
            "success": False,
            "message": "id is required"
        }), 400

    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        # -------------------------------------------------
        # ตรวจสอบข้อมูลที่เลือก
        # -------------------------------------------------
        cursor.execute("""
            SELECT *
            FROM sensor_data
            WHERE id = %s
              AND device_id = %s
            LIMIT 1
        """, (selected_id, device_id))

        selected = cursor.fetchone()

        if not selected:
            return jsonify({
                "success": False,
                "message": "Data not found"
            }), 404

        # -------------------------------------------------
        # ดึงข้อมูลก่อนหน้าประมาณ 30 รายการ
        # -------------------------------------------------
        cursor.execute("""
            SELECT *
            FROM sensor_data
            WHERE device_id = %s
              AND id <= %s
            ORDER BY id DESC
            LIMIT 30
        """, (device_id, selected_id))

        before_rows = cursor.fetchall()

        # -------------------------------------------------
        # ดึงข้อมูลหลังจากข้อมูลที่เลือกประมาณ 30 รายการ
        # -------------------------------------------------
        cursor.execute("""
            SELECT *
            FROM sensor_data
            WHERE device_id = %s
              AND id > %s
            ORDER BY id ASC
            LIMIT 30
        """, (device_id, selected_id))

        after_rows = cursor.fetchall()

        # -------------------------------------------------
        # รวมข้อมูล
        # -------------------------------------------------
        rows = list(reversed(before_rows)) + after_rows

        # -------------------------------------------------
        # แปลง timestamp ให้ JavaScript ใช้ง่าย
        # -------------------------------------------------
        for row in rows:
            if row.get("timestamp"):
                row["timestamp"] = row["timestamp"].isoformat()

        if selected.get("timestamp"):
            selected["timestamp"] = selected["timestamp"].isoformat()

        return jsonify({
            "success": True,
            "selected": selected,
            "data": rows
        })

    except Exception as e:
        print("GRAPH DATABASE ERROR:", e)

        return jsonify({
            "success": False,
            "message": "Database error"
        }), 500

    finally:
        cursor.close()
        conn.close()
# =========================================================
# EXPORT EXCEL
# =========================================================

@app.route(
    "/api/device/<device_id>/export/excel",
    methods=["GET"]
)
@login_required
@device_access_required
def export_excel(device_id):

    # -----------------------------------------
    # รับวันที่
    # เช่น 2026-10-05
    # -----------------------------------------

    date_value = request.args.get("date")


    conn = get_connection()

    cursor = conn.cursor(
        dictionary=True
    )


    try:

        # -----------------------------------------
        # ถ้ามีวันที่
        # -----------------------------------------

        if date_value:

            cursor.execute(
                """
                SELECT
                    timestamp,
                    accel_x,
                    accel_y,
                    accel_z,
                    pga,
                    peak_pga,
                    avg_pga,
                    pendulum,
                    level,
                    direction,
                    estimated_ml

                FROM sensor_data

                WHERE device_id = %s

                AND DATE(timestamp) = %s

                ORDER BY timestamp ASC
                """,
                (
                    device_id,
                    date_value
                )
            )


        # -----------------------------------------
        # ถ้าไม่มีวันที่
        # = เอาทั้งหมด
        # -----------------------------------------

        else:

            cursor.execute(
                """
                SELECT
                    timestamp,
                    accel_x,
                    accel_y,
                    accel_z,
                    pga,
                    peak_pga,
                    avg_pga,
                    pendulum,
                    level,
                    direction,
                    estimated_ml

                FROM sensor_data

                WHERE device_id = %s

                ORDER BY timestamp ASC
                """,
                (
                    device_id,
                )
            )


        rows = cursor.fetchall()


        # =================================================
        # CREATE EXCEL
        # =================================================

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Sensor Data"


        # -----------------------------------------
        # HEADER
        # -----------------------------------------

        headers = [

            "วันที่ / เวลา",

            "แกน X (G)",

            "แกน Y (G)",

            "แกน Z (G)",

            "PGA",

            "Peak PGA",

            "Average PGA",

            "Pendulum",

            "สถานะ",

            "ทิศทาง",

            "Estimated ML"

        ]


        sheet.append(headers)


        # -----------------------------------------
        # HEADER STYLE
        # -----------------------------------------

        for cell in sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )


        # -----------------------------------------
        # DATA
        # -----------------------------------------

        for row in rows:

            timestamp = row["timestamp"]


            if timestamp:

                timestamp = timestamp.strftime(
                    "%d/%m/%Y %H:%M:%S"
                )


            sheet.append([

                timestamp,

                row["accel_x"],

                row["accel_y"],

                row["accel_z"],

                row["pga"],

                row["peak_pga"],

                row["avg_pga"],

                row["pendulum"],

                row["level"],

                row["direction"],

                row["estimated_ml"]

            ])


        # -----------------------------------------
        # COLUMN WIDTH
        # -----------------------------------------

        widths = {

            "A": 22,

            "B": 15,

            "C": 15,

            "D": 15,

            "E": 15,

            "F": 15,

            "G": 15,

            "H": 15,

            "I": 15,

            "J": 15,

            "K": 18

        }


        for column, width in widths.items():

            sheet.column_dimensions[
                column
            ].width = width

        # =========================================
        # CREATE GRAPH SHEET
        # =========================================

        graph_sheet = workbook.create_sheet("กราฟ")

        graph_sheet["A1"] = "กราฟ PGA และ Pendulum"
        graph_sheet["A1"].font = Font(
            bold=True,
            size=16
        )

        # เพิ่มหัวตารางสำหรับข้อมูลที่ใช้ทำกราฟ
        graph_sheet.append([
            "วันที่ / เวลา",
            "PGA",
            "Pendulum"
        ])

        # ดึงข้อมูลจากชีต Sensor Data
        # A = วันที่ / เวลา
        # E = PGA
        # H = Pendulum
        for row_index in range(2, sheet.max_row + 1):
            graph_sheet.append([
                sheet.cell(row=row_index, column=1).value,
                sheet.cell(row=row_index, column=5).value or 0,
                sheet.cell(row=row_index, column=8).value or 0
            ])

        # สร้างกราฟเส้น
        chart = LineChart()
        chart.title = "PGA และ Pendulum ตามเวลา"
        chart.style = 13
        chart.y_axis.title = "ค่า"
        chart.x_axis.title = "วันที่ / เวลา"
        chart.height = 12
        chart.width = 24
        chart.display_blanks = "gap"

        # PGA และ Pendulum
        data = Reference(
            graph_sheet,
            min_col=2,
            max_col=3,
            min_row=2,
            max_row=graph_sheet.max_row
        )

        # ใช้เวลาเป็นแกน X
        categories = Reference(
            graph_sheet,
            min_col=1,
            min_row=3,
            max_row=graph_sheet.max_row
        )

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        # วางกราฟในชีต
        graph_sheet.add_chart(chart, "E2")

        # ปรับความกว้างคอลัมน์
        graph_sheet.column_dimensions["A"].width = 24
        graph_sheet.column_dimensions["B"].width = 16
        graph_sheet.column_dimensions["C"].width = 16
        # -----------------------------------------
        # CREATE FILE IN MEMORY
        # -----------------------------------------

        output = BytesIO()

        workbook.save(output)

        output.seek(0)


        # -----------------------------------------
        # FILE NAME
        # -----------------------------------------

        filename = (
            f"{device_id}_sensor_data"
        )


        if date_value:

            filename += (
                f"_{date_value}"
            )


        filename += ".xlsx"


        # -----------------------------------------
        # DOWNLOAD
        # -----------------------------------------

        return send_file(

            output,

            as_attachment=True,

            download_name=filename,

            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )

        )


    except Exception as e:

        print("EXCEL ERROR:", e)

        return jsonify({

            "success": False,

            "message": "Excel export error"

        }), 500


    finally:

        cursor.close()
        conn.close()

# =========================================================
# EXPORT PDF
# =========================================================

@app.route(
    "/api/device/<device_id>/export/pdf",
    methods=["GET"]
)
@login_required
@device_access_required
def export_pdf(device_id):

    # -----------------------------------------
    # รับวันที่
    # -----------------------------------------

    date_value = request.args.get("date")


    
    period = request.args.get("period", "all")
    date_value = request.args.get("date")
    month_value = request.args.get("month")
    year_value = request.args.get("year")

    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

        



    try:

        # -----------------------------------------
        # ถ้ามีวันที่
        # -----------------------------------------

        if date_value:

            cursor.execute(
                """
                SELECT
                    timestamp,
                    accel_x,
                    accel_y,
                    accel_z,
                    pga,
                    peak_pga,
                    avg_pga,
                    pendulum,
                    level,
                    direction,
                    estimated_ml

                FROM sensor_data

                WHERE device_id = %s

                AND DATE(timestamp) = %s

                ORDER BY timestamp ASC
                """,
                (
                    device_id,
                    date_value
                )
            )


        # -----------------------------------------
        # ไม่มีวันที่
        # = เอาทั้งหมด
        # -----------------------------------------

        else:

            cursor.execute(
                """
                SELECT
                    timestamp,
                    accel_x,
                    accel_y,
                    accel_z,
                    pga,
                    peak_pga,
                    avg_pga,
                    pendulum,
                    level,
                    direction,
                    estimated_ml

                FROM sensor_data

                WHERE device_id = %s

                ORDER BY timestamp ASC
                """,
                (
                    device_id,
                )
            )


        rows = cursor.fetchall()


        # =================================================
        # CREATE PDF
        # =================================================

        output = BytesIO()


        document = SimpleDocTemplate(

            output,

            pagesize=landscape(A4),

            rightMargin=20,

            leftMargin=20,

            topMargin=20,

            bottomMargin=20

        )


        styles = getSampleStyleSheet()


        title_style = styles["Title"]

        title_style.alignment = TA_CENTER


        elements = []


        # -----------------------------------------
        # TITLE
        # -----------------------------------------

        title = (
            "ESP32 Sensor Monitoring Report"
        )


        if date_value:

            title += (
                f" - {date_value}"
            )


        elements.append(

            Paragraph(
                title,
                title_style
            )

        )


        elements.append(
            Spacer(1, 15)
        )


        # =================================================
        # TABLE
        # =================================================

        table_data = [

            [

                "Date / Time",

                "X (G)",

                "Y (G)",

                "Z (G)",

                "PGA",

                "Peak PGA",

                "Avg PGA",

                "Pendulum",

                "Level",

                "Direction",

                "ML"

            ]

        ]


        # -----------------------------------------
        # DATA
        # -----------------------------------------

        for row in rows:

            timestamp = row["timestamp"]


            if timestamp:

                timestamp = timestamp.strftime(
                    "%d/%m/%Y %H:%M:%S"
                )

            else:

                timestamp = "-"


            table_data.append([

                timestamp,

                f"{float(row['accel_x']):.4f}",

                f"{float(row['accel_y']):.4f}",

                f"{float(row['accel_z']):.4f}",

                f"{float(row['pga']):.4f}",

                f"{float(row['peak_pga']):.4f}",

                f"{float(row['avg_pga']):.4f}",

                f"{float(row['pendulum']):.2f}",

                str(row["level"]),

                str(row["direction"]),

                f"{float(row['estimated_ml']):.2f}"

            ])


        # -----------------------------------------
        # TABLE
        # -----------------------------------------

        table = Table(
            table_data,
            repeatRows=1
        )


        table.setStyle(

            TableStyle([

                (

                    "BACKGROUND",

                    (0, 0),

                    (-1, 0),

                    colors.HexColor(
                        "#16304f"
                    )

                ),

                (

                    "TEXTCOLOR",

                    (0, 0),

                    (-1, 0),

                    colors.white

                ),

                (

                    "FONTNAME",

                    (0, 0),

                    (-1, 0),

                    "Helvetica-Bold"

                ),

                (

                    "ALIGN",

                    (0, 0),

                    (-1, -1),

                    "CENTER"

                ),

                (

                    "VALIGN",

                    (0, 0),

                    (-1, -1),

                    "MIDDLE"

                ),

                (

                    "GRID",

                    (0, 0),

                    (-1, -1),

                    0.5,

                    colors.grey

                ),

                (

                    "FONTSIZE",

                    (0, 0),

                    (-1, -1),

                    7

                ),

                (

                    "ROWBACKGROUNDS",

                    (0, 1),

                    (-1, -1),

                    [

                        colors.white,

                        colors.HexColor(
                            "#f2f2f2"
                        )

                    ]

                )

            ])

        )



        elements.append(table)

        # =================================================
        # CREATE PGA + PENDULUM GRAPH FOR PDF
        # =================================================

        if rows:
            timestamps = [
                row["timestamp"].strftime("%d/%m %H:%M")
                if row["timestamp"] else "-"
                for row in rows
            ]

            pga_values = [
                float(row["pga"] or 0)
                for row in rows
            ]

            pendulum_values = [
                float(row["pendulum"] or 0)
                for row in rows
            ]

            fig, ax = plt.subplots(figsize=(11, 4.5))

            ax.plot(
                range(len(rows)),
                pga_values,
                label="PGA",
                marker=".",
                linewidth=1.5
            )

            ax.plot(
                range(len(rows)),
                pendulum_values,
                label="Pendulum",
                marker=".",
                linewidth=1.5
            )

            ax.set_title("PGA and Pendulum History")
            ax.set_xlabel("Date / Time")
            ax.set_ylabel("Value")
            ax.grid(True, alpha=0.3)
            ax.legend()

            # ลดจำนวนป้ายเวลา เพื่อไม่ให้ทับกัน
            step = max(1, len(timestamps) // 10)

            ax.set_xticks(range(0, len(timestamps), step))
            ax.set_xticklabels(
                [timestamps[i] for i in range(0, len(timestamps), step)],
                rotation=35,
                ha="right"
            )

            fig.tight_layout()

            # บันทึกกราฟไว้ในหน่วยความจำ
            graph_buffer = BytesIO()
            fig.savefig(
                graph_buffer,
                format="png",
                dpi=160,
                bbox_inches="tight"
            )
            plt.close(fig)
            graph_buffer.seek(0)

            # ขึ้นหน้าใหม่ก่อนแสดงกราฟ
            from reportlab.platypus import PageBreak

            elements.append(PageBreak())
            elements.append(
                Paragraph("PGA and Pendulum Graph", title_style)
            )
            elements.append(Spacer(1, 12))
            elements.append(
                RLImage(graph_buffer, width=750, height=300)
            )

# BUILD PDF
        document.build(elements)


        output.seek(0)


        # -----------------------------------------
        # FILE NAME
        # -----------------------------------------

        filename = (
            f"{device_id}_sensor_data"
        )


        if date_value:

            filename += (
                f"_{date_value}"
            )


        filename += ".pdf"


        # -----------------------------------------
        # DOWNLOAD
        # -----------------------------------------

        return send_file(

            output,

            as_attachment=True,

            download_name=filename,

            mimetype="application/pdf"

        )


    except Exception as e:

        print("PDF ERROR:", e)

        return jsonify({

            "success": False,

            "message": "PDF export error"

        }), 500


    finally:

        cursor.close()
        conn.close()
        




def get_export_rows(cursor, device_id):
    period = request.args.get("period", "all")
    date_value = request.args.get("date")
    month_value = request.args.get("month")
    year_value = request.args.get("year")

    query = """
        SELECT
            timestamp, accel_x, accel_y, accel_z,
            pga, peak_pga, avg_pga, pendulum,
            level, direction, estimated_ml
        FROM sensor_data
        WHERE device_id = %s
    """
    params = [device_id]

    if period == "day":
        try:
            datetime.strptime(date_value or "", "%Y-%m-%d")
        except ValueError:
            return None, ("รูปแบบวันที่ไม่ถูกต้อง", 400)

        query += """
            AND timestamp >= %s
            AND timestamp < DATE_ADD(%s, INTERVAL 1 DAY)
        """
        params.extend([date_value, date_value])

    elif period == "month":
        try:
            datetime.strptime(month_value or "", "%Y-%m")
        except ValueError:
            return None, ("รูปแบบเดือนไม่ถูกต้อง", 400)

        query += """
            AND timestamp >= STR_TO_DATE(
                CONCAT(%s, '-01'), '%%Y-%%m-%%d'
            )
            AND timestamp < DATE_ADD(
                STR_TO_DATE(
                    CONCAT(%s, '-01'), '%%Y-%%m-%%d'
                ),
                INTERVAL 1 MONTH
            )
        """
        params.extend([month_value, month_value])

    elif period == "year":
        if (
            not year_value
            or not year_value.isdigit()
            or not 2000 <= int(year_value) <= 2100
        ):
            return None, ("รูปแบบปีไม่ถูกต้อง", 400)

        query += """
            AND timestamp >= %s
            AND timestamp < %s
        """
        params.extend([
            f"{year_value}-01-01",
            f"{int(year_value) + 1}-01-01"
        ])

    elif period != "all":
        return None, ("ช่วงข้อมูลไม่ถูกต้อง", 400)

    query += " ORDER BY timestamp ASC"
    cursor.execute(query, tuple(params))

    return cursor.fetchall(), None

# =========================================================
# GET DEVICES
# =========================================================


@app.route("/api/devices", methods=["GET"])
@login_required
def get_devices():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    try:
        cursor.execute("""
            SELECT
                d.device_id,
                d.status,
                d.ip_address,
                d.last_seen,
                ud.added_at
            FROM user_devices AS ud
            JOIN devices AS d
                ON d.device_id = ud.device_id
            WHERE ud.user_id = %s
            ORDER BY ud.added_at DESC
        """, (session["user_id"],))

        rows = cursor.fetchall()

        for row in rows:
            for key in ("last_seen", "added_at"):
                if row.get(key):
                    row[key] = row[key].isoformat()

        return jsonify({
            "success": True,
            "count": len(rows),
            "devices": rows
        })

    finally:
        cursor.close()
        conn.close()

# =========================================================
# RUN SERVER
# =========================================================

if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )


    app.run(

        host="0.0.0.0",

        port=port,

        debug=False

    )